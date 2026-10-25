#!/usr/bin/env python3
"""
inject_to_excel.py
Inject exported CSV span records into the Excel report template (trace_report_template.xlsx).
Creates a rich AWS X-Ray style report with:
- Executive Summary & Multi-Trace Index
- Hierarchical Span Waterfall with Tree Visualization
- In-Cell Timeline Gantt Bars and Native Stacked Bar Chart
- Full Raw Data View
Designed for RHEL 9.8 / Python 3.9+ with openpyxl.
"""

import sys
import os
import csv
import argparse
from datetime import datetime, timezone
from collections import defaultdict
from typing import List, Dict, Any, Tuple, Optional

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, Reference, Series

# Palette matching AWS X-Ray / CloudWatch
NAVY_HEADER = "1B2A4A"
WHITE = "FFFFFF"
GRAY_TEXT = "5A6A85"
CARD_BG = "F8FAFC"
CARD_BORDER = "E2E8F0"
BORDER_LIGHT = "E5E7EB"
BORDER_DARK = "CBD5E1"

# Status colors
SUCCESS_BG = "D1E7DD"
SUCCESS_FG = "0F5132"
SUCCESS_BAR = "34A853"   # Green

WARNING_BG = "FFF3CD"
WARNING_FG = "664D03"
WARNING_BAR = "FBBC05"   # Amber / Orange

DANGER_BG = "F8D7DA"
DANGER_FG = "842029"
DANGER_BAR = "EA4335"    # Crimson Red

CLIENT_BAR = "4285F4"    # Blue for Client / DB
INTERNAL_BAR = "9AA0A6"  # Gray for Internal

TIMELINE_BG = "F1F5F9"
ROW_ALT = "F8FAFC"
ROW_WHITE = "FFFFFF"
FONT_FAMILY = "Segoe UI"


def create_thin_border(color=BORDER_LIGHT):
    side = Side(border_style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def parse_timestamp_to_epoch_ms(ts_str: str) -> float:
    """Parse timestamp string into epoch milliseconds."""
    if not ts_str:
        return 0.0
    s = ts_str.strip()
    # Normalize ISO format (e.g. 2026-10-26T04:30:00.000Z -> 2026-10-26 04:30:00.000)
    s = s.replace("T", " ").replace("Z", "")
    for fmt in (
        "%Y-%m-%d %H:%M:%S.%f",
        "%Y-%m-%d %H:%M:%S",
        "%Y-%m-%d %H:%M",
        "%Y/%m/%d %H:%M:%S.%f",
        "%Y/%m/%d %H:%M:%S",
    ):
        try:
            dt = datetime.strptime(s, fmt)
            return dt.timestamp() * 1000.0
        except ValueError:
            pass
    try:
        # Fallback if already numeric
        return float(s)
    except ValueError:
        return 0.0


class SpanNode:
    """Represents a span node in the execution hierarchy tree."""
    def __init__(self, data: Dict[str, Any]):
        self.data = data
        self.span_id = data.get("span_id", "")
        self.parent_span_id = data.get("parent_span_id", "")
        self.trace_id = data.get("trace_id", "")
        self.service_name = data.get("service_name", "UnknownService")
        self.operation_name = data.get("operation_name", "UnknownOperation")
        self.http_status = str(data.get("http_status", "")).strip()
        self.status_code = str(data.get("status_code", "")).strip().upper()
        self.kind = str(data.get("kind", "")).strip().upper()
        self.timestamp_str = data.get("timestamp", "")
        self.epoch_ms = parse_timestamp_to_epoch_ms(self.timestamp_str)

        try:
            self.duration_ms = float(data.get("duration_ms", 0.0))
        except (ValueError, TypeError):
            self.duration_ms = 0.0

        self.children: List["SpanNode"] = []
        self.depth = 0
        self.offset_ms = 0.0
        self.tree_prefix = ""


def build_trace_trees(rows: List[Dict[str, Any]]) -> Dict[str, Dict[str, Any]]:
    """
    Group spans by trace_id and build hierarchical execution tree.
    Returns dictionary of trace summaries and root span trees.
    """
    traces_dict = defaultdict(list)
    for r in rows:
        trace_id = r.get("trace_id", "")
        if trace_id:
            traces_dict[trace_id].append(r)

    results = {}

    for trace_id, span_rows in traces_dict.items():
        # Create SpanNode objects
        nodes_by_id = {r["span_id"]: SpanNode(r) for r in span_rows if r.get("span_id")}
        
        # Determine root spans (parent_span_id empty or parent not in this trace)
        roots: List[SpanNode] = []
        for node in nodes_by_id.values():
            if not node.parent_span_id or node.parent_span_id not in nodes_by_id:
                roots.append(node)
            else:
                nodes_by_id[node.parent_span_id].children.append(node)

        # Sort children by epoch_ms / start time
        for node in nodes_by_id.values():
            node.children.sort(key=lambda n: n.epoch_ms)

        roots.sort(key=lambda n: n.epoch_ms)

        # Find min timestamp for trace start
        all_nodes = list(nodes_by_id.values())
        if all_nodes:
            trace_start_epoch_ms = min(n.epoch_ms for n in all_nodes)
        else:
            trace_start_epoch_ms = 0.0

        # Calculate offsets
        for node in all_nodes:
            node.offset_ms = max(0.0, node.epoch_ms - trace_start_epoch_ms)

        # Total trace duration
        if all_nodes:
            trace_end_ms = max(n.offset_ms + n.duration_ms for n in all_nodes)
            total_duration_ms = max(trace_end_ms, 1.0)
        else:
            total_duration_ms = 1.0

        # Flatten hierarchy via preorder DFS
        ordered_spans: List[SpanNode] = []

        def traverse(node: SpanNode, depth: int, is_last: bool, prefix_acc: str):
            node.depth = depth
            branch = "└── " if is_last else "├── "
            if depth == 0:
                node.tree_prefix = ""
            else:
                node.tree_prefix = prefix_acc + branch

            ordered_spans.append(node)

            child_prefix = prefix_acc + ("    " if is_last else "│   ") if depth > 0 else ""
            for idx, child in enumerate(node.children):
                last_child = (idx == len(node.children) - 1)
                traverse(child, depth + 1, last_child, child_prefix)

        for idx, root in enumerate(roots):
            traverse(root, depth=0, is_last=(idx == len(roots) - 1), prefix_acc="")

        # Compute trace status
        has_5xx = any(
            (n.http_status.startswith("5") or n.status_code == "ERROR")
            for n in all_nodes
        )
        has_4xx = any(n.http_status.startswith("4") for n in all_nodes)

        if has_5xx:
            status_label = "FAULT"
        elif has_4xx:
            status_label = "ERROR"
        else:
            status_label = "OK"

        root_service = roots[0].service_name if roots else "Unknown"
        root_op = roots[0].operation_name if roots else "Unknown"
        start_time_str = roots[0].timestamp_str if roots else ""
        services_set = {n.service_name for n in all_nodes}

        results[trace_id] = {
            "trace_id": trace_id,
            "root_service": root_service,
            "root_op": root_op,
            "start_time": start_time_str,
            "total_duration_ms": total_duration_ms,
            "spans_count": len(all_nodes),
            "services_count": len(services_set),
            "services": sorted(list(services_set)),
            "status": status_label,
            "error_count": sum(1 for n in all_nodes if n.status_code == "ERROR" or n.http_status.startswith("4") or n.http_status.startswith("5")),
            "roots": roots,
            "ordered_spans": ordered_spans,
        }

    return results


def get_status_styling(status_str: str, http_status: str = "", kind: str = ""):
    """Return background and foreground colors based on span status."""
    st = str(status_str).upper()
    hs = str(http_status).strip()

    if hs.startswith("5") or st in ("FAULT", "ERROR"):
        return DANGER_BG, DANGER_FG, DANGER_BAR, "FAULT" if hs.startswith("5") else "ERROR"
    elif hs.startswith("4"):
        return WARNING_BG, WARNING_FG, WARNING_BAR, "ERROR"
    else:
        bar_color = CLIENT_BAR if kind == "CLIENT" else SUCCESS_BAR
        return SUCCESS_BG, SUCCESS_FG, bar_color, "OK"


def populate_summary_sheet(ws, traces_dict: Dict[str, Dict[str, Any]], metadata: Dict[str, str]):
    """Fill in the Summary sheet with metadata, KPI cards, and traces table."""
    total_traces = len(traces_dict)
    total_spans = sum(t["spans_count"] for t in traces_dict.values())
    ok_traces = sum(1 for t in traces_dict.values() if t["status"] == "OK")
    error_traces = total_traces - ok_traces
    max_duration = max((t["total_duration_ms"] for t in traces_dict.values()), default=0.0)

    # Fill metadata
    ws["C4"].value = metadata.get("log_group", "aws/spans")
    ws["G4"].value = f"{metadata.get('start_time', 'N/A')}  ~  {metadata.get('end_time', 'N/A')}"
    ws["C5"].value = metadata.get("generated_at", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"))
    ws["G5"].value = metadata.get("filter_criteria", "All Transactions")

    # Fill KPI Cards
    ws["B8"].value = total_traces
    ws["E8"].value = total_spans
    ws["H8"].value = ok_traces
    ws["J8"].value = error_traces
    ws["L8"].value = f"{max_duration:,.1f} ms"

    # Fill Traces Table
    start_row = 13
    for idx, (trace_id, t_info) in enumerate(traces_dict.items()):
        row = start_row + idx
        bg_col = ROW_WHITE if idx % 2 == 0 else ROW_ALT
        fill_standard = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")

        status_bg, status_fg, _, status_text = get_status_styling(t_info["status"])

        # Col B: Trace ID
        c_b = ws.cell(row=row, column=2, value=trace_id)
        c_b.font = Font(name=FONT_FAMILY, size=9)
        c_b.alignment = Alignment(horizontal="left", vertical="center")
        c_b.fill = fill_standard
        c_b.border = create_thin_border()

        # Col C: Start Time
        c_c = ws.cell(row=row, column=3, value=t_info["start_time"])
        c_c.font = Font(name=FONT_FAMILY, size=9)
        c_c.alignment = Alignment(horizontal="center", vertical="center")
        c_c.fill = fill_standard
        c_c.border = create_thin_border()

        # Col D: Root Service
        c_d = ws.cell(row=row, column=4, value=t_info["root_service"])
        c_d.font = Font(name=FONT_FAMILY, size=9, bold=True)
        c_d.alignment = Alignment(horizontal="left", vertical="center")
        c_d.fill = fill_standard
        c_d.border = create_thin_border()

        # Col E: Root Operation
        c_e = ws.cell(row=row, column=5, value=t_info["root_op"])
        c_e.font = Font(name=FONT_FAMILY, size=9)
        c_e.alignment = Alignment(horizontal="left", vertical="center")
        c_e.fill = fill_standard
        c_e.border = create_thin_border()

        # Col F: Duration
        c_f = ws.cell(row=row, column=6, value=round(t_info["total_duration_ms"], 2))
        c_f.font = Font(name=FONT_FAMILY, size=9)
        c_f.number_format = '#,##0.00 "ms"'
        c_f.alignment = Alignment(horizontal="right", vertical="center")
        c_f.fill = fill_standard
        c_f.border = create_thin_border()

        # Col G: Spans Count
        c_g = ws.cell(row=row, column=7, value=t_info["spans_count"])
        c_g.font = Font(name=FONT_FAMILY, size=9)
        c_g.alignment = Alignment(horizontal="center", vertical="center")
        c_g.fill = fill_standard
        c_g.border = create_thin_border()

        # Col H: Services Count
        c_h = ws.cell(row=row, column=8, value=t_info["services_count"])
        c_h.font = Font(name=FONT_FAMILY, size=9)
        c_h.alignment = Alignment(horizontal="center", vertical="center")
        c_h.fill = fill_standard
        c_h.border = create_thin_border()

        # Col I: Status Badge
        c_i = ws.cell(row=row, column=9, value=status_text)
        c_i.font = Font(name=FONT_FAMILY, size=9, bold=True, color=status_fg)
        c_i.alignment = Alignment(horizontal="center", vertical="center")
        c_i.fill = PatternFill(start_color=status_bg, end_color=status_bg, fill_type="solid")
        c_i.border = create_thin_border()

        # Col J: Error Count
        c_j = ws.cell(row=row, column=10, value=t_info["error_count"])
        c_j.font = Font(name=FONT_FAMILY, size=9, bold=(t_info["error_count"] > 0),
                        color=DANGER_FG if t_info["error_count"] > 0 else GRAY_TEXT)
        c_j.alignment = Alignment(horizontal="center", vertical="center")
        c_j.fill = fill_standard
        c_j.border = create_thin_border()

        # Col K: Link to Waterfall Sheet
        target_sheet = "Trace_Waterfall" if idx == 0 else f"Trace_{idx + 1}"
        c_k = ws.cell(row=row, column=11, value="View Timeline →")
        c_k.font = Font(name=FONT_FAMILY, size=9, bold=True, color="0D6EFD", underline="single")
        c_k.hyperlink = f"#'{target_sheet}'!A1"
        c_k.alignment = Alignment(horizontal="center", vertical="center")
        c_k.fill = fill_standard
        c_k.border = create_thin_border()

        ws.row_dimensions[row].height = 20


def populate_waterfall_sheet(ws, trace_info: Dict[str, Any], sheet_title: str = "Trace_Waterfall"):
    """Populate a Waterfall sheet for a specific trace with tree and timeline."""
    ws.title = sheet_title

    # Overview KPI Cards
    ws["B6"].value = trace_info["trace_id"]
    status_bg, status_fg, _, status_text = get_status_styling(trace_info["status"])
    ws["F6"].value = status_text
    ws["F6"].font = Font(name=FONT_FAMILY, size=14, bold=True, color=status_fg)
    ws["F6"].fill = PatternFill(start_color=status_bg, end_color=status_bg, fill_type="solid")

    root_str = f"{trace_info['root_service']} : {trace_info['root_op']}"
    ws["H6"].value = root_str
    ws["M6"].value = f"{trace_info['total_duration_ms']:,.2f} ms"
    ws["O6"].value = trace_info["spans_count"]
    ws["Q6"].value = trace_info["start_time"]

    # Table rows
    start_row = 11
    ordered_spans: List[SpanNode] = trace_info["ordered_spans"]
    total_dur = max(trace_info["total_duration_ms"], 1.0)
    timeline_cols_count = 20
    timeline_start_col = 10  # Col J

    for idx, span in enumerate(ordered_spans):
        row = start_row + idx
        bg_col = ROW_WHITE if idx % 2 == 0 else ROW_ALT
        fill_standard = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")

        s_bg, s_fg, s_bar, s_label = get_status_styling(span.status_code, span.http_status, span.kind)

        # Col B: Status
        c_status = ws.cell(row=row, column=2, value=s_label)
        c_status.font = Font(name=FONT_FAMILY, size=8, bold=True, color=s_fg)
        c_status.alignment = Alignment(horizontal="center", vertical="center")
        c_status.fill = PatternFill(start_color=s_bg, end_color=s_bg, fill_type="solid")
        c_status.border = create_thin_border()

        # Col C: Hierarchy & Operation
        display_name = f"{span.tree_prefix}{span.operation_name}"
        c_name = ws.cell(row=row, column=3, value=display_name)
        c_name.font = Font(name="Consolas" if span.tree_prefix else FONT_FAMILY,
                           size=9, bold=(span.depth == 0))
        c_name.alignment = Alignment(horizontal="left", vertical="center")
        c_name.fill = fill_standard
        c_name.border = create_thin_border()

        # Col D: Service
        c_svc = ws.cell(row=row, column=4, value=span.service_name)
        c_svc.font = Font(name=FONT_FAMILY, size=9, bold=(span.depth <= 1))
        c_svc.alignment = Alignment(horizontal="left", vertical="center")
        c_svc.fill = fill_standard
        c_svc.border = create_thin_border()

        # Col E: Kind
        c_kind = ws.cell(row=row, column=5, value=span.kind or "INTERNAL")
        c_kind.font = Font(name=FONT_FAMILY, size=8, color=GRAY_TEXT)
        c_kind.alignment = Alignment(horizontal="center", vertical="center")
        c_kind.fill = fill_standard
        c_kind.border = create_thin_border()

        # Col F: HTTP Status
        c_http = ws.cell(row=row, column=6, value=span.http_status)
        c_http.font = Font(name=FONT_FAMILY, size=9)
        c_http.alignment = Alignment(horizontal="center", vertical="center")
        c_http.fill = fill_standard
        c_http.border = create_thin_border()

        # Col G: Start Offset
        c_off = ws.cell(row=row, column=7, value=round(span.offset_ms, 2))
        c_off.font = Font(name=FONT_FAMILY, size=9)
        c_off.number_format = '"+ "#,##0.00 "ms"'
        c_off.alignment = Alignment(horizontal="right", vertical="center")
        c_off.fill = fill_standard
        c_off.border = create_thin_border()

        # Col H: Duration
        c_dur = ws.cell(row=row, column=8, value=round(span.duration_ms, 2))
        c_dur.font = Font(name=FONT_FAMILY, size=9, bold=True)
        c_dur.number_format = '#,##0.00 "ms"'
        c_dur.alignment = Alignment(horizontal="right", vertical="center")
        c_dur.fill = fill_standard
        c_dur.border = create_thin_border()

        # Col I: % of Trace
        pct_share = span.duration_ms / total_dur
        c_pct = ws.cell(row=row, column=9, value=pct_share)
        c_pct.font = Font(name=FONT_FAMILY, size=9)
        c_pct.number_format = '0.0%'
        c_pct.alignment = Alignment(horizontal="right", vertical="center")
        c_pct.fill = fill_standard
        c_pct.border = create_thin_border()

        # Timeline visualizer (Cols J to AC: 20 buckets)
        start_ratio = span.offset_ms / total_dur
        end_ratio = (span.offset_ms + span.duration_ms) / total_dur

        bucket_start = min(timeline_cols_count - 1, max(0, int(start_ratio * timeline_cols_count)))
        bucket_end = min(timeline_cols_count - 1, max(bucket_start, int(end_ratio * timeline_cols_count)))

        bar_fill = PatternFill(start_color=s_bar, end_color=s_bar, fill_type="solid")
        empty_fill = PatternFill(start_color=TIMELINE_BG, end_color=TIMELINE_BG, fill_type="solid")

        for b_idx in range(timeline_cols_count):
            col_idx = timeline_start_col + b_idx
            c_bar = ws.cell(row=row, column=col_idx)
            c_bar.border = create_thin_border(BORDER_LIGHT)
            if bucket_start <= b_idx <= bucket_end:
                c_bar.fill = bar_fill
            else:
                c_bar.fill = empty_fill

        ws.row_dimensions[row].height = 20

    # Add Native Stacked Bar Chart for Gantt visualization
    if ordered_spans:
        chart = BarChart()
        chart.type = "bar"
        chart.grouping = "stacked"
        chart.overlap = 100
        chart.title = f"Trace Gantt Waterfall: {trace_info['root_op']}"
        chart.style = 10
        chart.height = min(18, max(8, len(ordered_spans) * 0.75 + 3))
        chart.width = 22

        # Data: Offset (Col G) and Duration (Col H)
        data = Reference(ws, min_col=7, min_row=10, max_col=8, max_row=start_row + len(ordered_spans) - 1)
        cats = Reference(ws, min_col=3, min_row=start_row, max_row=start_row + len(ordered_spans) - 1)
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(cats)

        # Invert vertical axis so first span appears at top
        chart.y_axis.scaling.orientation = "maxMin"
        chart.x_axis.title = "Milliseconds (ms)"

        # Set series 1 (Offset) invisible
        if len(chart.series) >= 2:
            chart.series[0].graphicalProperties.noFill = True
            chart.series[0].graphicalProperties.line.noFill = True

        chart_pos_row = start_row + len(ordered_spans) + 2
        ws.add_chart(chart, f"B{chart_pos_row}")


def populate_raw_sheet(ws, raw_rows: List[Dict[str, Any]]):
    """Fill raw spans data sheet."""
    start_row = 5
    for idx, r in enumerate(raw_rows):
        row = start_row + idx
        bg_col = ROW_WHITE if idx % 2 == 0 else ROW_ALT
        fill_standard = PatternFill(start_color=bg_col, end_color=bg_col, fill_type="solid")

        vals = [
            r.get("timestamp", ""),
            r.get("trace_id", ""),
            r.get("span_id", ""),
            r.get("parent_span_id", ""),
            r.get("service_name", ""),
            r.get("operation_name", ""),
            float(r["duration_ms"]) if r.get("duration_ms") else "",
            int(r["duration_nano"]) if r.get("duration_nano") else "",
            r.get("http_status", ""),
            r.get("http_method", ""),
            r.get("http_url", ""),
            r.get("status_code", ""),
            r.get("status_message", ""),
            r.get("kind", ""),
        ]

        for col_idx, val in enumerate(vals, start=1):
            cell = ws.cell(row=row, column=col_idx, value=val)
            cell.font = Font(name=FONT_FAMILY, size=9)
            cell.fill = fill_standard
            cell.border = create_thin_border()
            if col_idx == 7 and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = Alignment(horizontal="right", vertical="center")
            elif col_idx in (8, 9):
                cell.alignment = Alignment(horizontal="center", vertical="center")
            else:
                cell.alignment = Alignment(horizontal="left", vertical="center")

        ws.row_dimensions[row].height = 18


def inject_report(
    csv_path: str,
    template_path: str,
    output_path: str,
    target_trace_id: Optional[str] = None,
    start_time: str = "",
    end_time: str = "",
    filter_criteria: str = "",
    log_group: str = "aws/spans",
):
    """Main injection function reading CSV and populating Excel workbook."""
    if not os.path.exists(csv_path):
        raise FileNotFoundError(f"CSV file not found: {csv_path}")

    # Read CSV records
    raw_rows = []
    with open(csv_path, "r", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        for r in reader:
            raw_rows.append(r)

    # Reconstruct trace trees
    traces_dict = build_trace_trees(raw_rows)

    if not os.path.exists(template_path):
        # Auto-generate template if missing
        from generate_template import create_template
        create_template(template_path)

    wb = openpyxl.load_workbook(template_path)

    metadata = {
        "log_group": log_group,
        "start_time": start_time or (raw_rows[0]["timestamp"] if raw_rows else "N/A"),
        "end_time": end_time or (raw_rows[-1]["timestamp"] if raw_rows else "N/A"),
        "generated_at": datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M:%S UTC"),
        "filter_criteria": filter_criteria or (f"traceId={target_trace_id}" if target_trace_id else "All Spans"),
    }

    # 1. Populate Summary Sheet
    if "Summary" in wb.sheetnames:
        populate_summary_sheet(wb["Summary"], traces_dict, metadata)

    # 2. Select / Populate Trace Waterfall
    if traces_dict:
        # Determine primary trace
        if target_trace_id and target_trace_id in traces_dict:
            primary_trace_id = target_trace_id
        else:
            # Pick trace with error, or longest trace
            sorted_traces = sorted(
                traces_dict.values(),
                key=lambda t: (1 if t["status"] != "OK" else 0, t["total_duration_ms"]),
                reverse=True,
            )
            primary_trace_id = sorted_traces[0]["trace_id"]

        primary_trace = traces_dict[primary_trace_id]
        if "Trace_Waterfall" in wb.sheetnames:
            populate_waterfall_sheet(wb["Trace_Waterfall"], primary_trace, sheet_title="Trace_Waterfall")

        # For additional traces (up to top 5), copy template or create extra sheets
        other_traces = [t for t in traces_dict.values() if t["trace_id"] != primary_trace_id][:4]
        for idx, other_t in enumerate(other_traces, start=2):
            new_sheet_name = f"Trace_{idx}"
            if "Trace_Waterfall" in wb.sheetnames:
                new_ws = wb.copy_worksheet(wb["Trace_Waterfall"])
                populate_waterfall_sheet(new_ws, other_t, sheet_title=new_sheet_name)
    else:
        # No traces found
        if "Trace_Waterfall" in wb.sheetnames:
            ws_wat = wb["Trace_Waterfall"]
            ws_wat["B6"].value = "No traces matched the query criteria."

    # 3. Populate Raw_Data Sheet
    if "Raw_Data" in wb.sheetnames:
        populate_raw_sheet(wb["Raw_Data"], raw_rows)

    os.makedirs(os.path.dirname(os.path.abspath(output_path)), exist_ok=True)
    wb.save(output_path)
    print(f"Generated report successfully: {output_path}")
    print(f"Total Traces: {len(traces_dict)}, Total Spans: {len(raw_rows)}")


def main():
    parser = argparse.ArgumentParser(description="Inject CSV spans into Excel Report Template.")
    parser.add_argument("--csv", required=True, help="Path to input CSV file")
    parser.add_argument("--template", default="templates/trace_report_template.xlsx", help="Path to Excel template")
    parser.add_argument("--output", required=True, help="Path to output Excel report")
    parser.add_argument("--trace-id", default=None, help="Specific Trace ID to focus on")
    parser.add_argument("--start-time", default="", help="Query start time")
    parser.add_argument("--end-time", default="", help="Query end time")
    parser.add_argument("--filter", default="", help="Filter criteria string")
    parser.add_argument("--log-group", default="aws/spans", help="CloudWatch Logs log group")

    args = parser.parse_args()

    inject_report(
        csv_path=args.csv,
        template_path=args.template,
        output_path=args.output,
        target_trace_id=args.trace_id,
        start_time=args.start_time,
        end_time=args.end_time,
        filter_criteria=args.filter,
        log_group=args.log_group,
    )


if __name__ == "__main__":
    main()
