#!/usr/bin/env python3
"""
test_inject_to_excel.py
Unit and integration tests for inject_to_excel.py.
"""

import unittest
import tempfile
import os
import sys
import openpyxl

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "scripts")))
from inject_to_excel import (
    parse_timestamp_to_epoch_ms,
    build_trace_trees,
    inject_report,
    SpanNode,
)
from generate_template import create_template


class TestInjectToExcel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Create a temporary master template
        cls.template_file = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        cls.template_path = cls.template_file.name
        cls.template_file.close()
        create_template(cls.template_path)

    @classmethod
    def tearDownClass(cls):
        if os.path.exists(cls.template_path):
            os.remove(cls.template_path)

    def test_parse_timestamp_to_epoch_ms(self):
        # ISO format
        ms1 = parse_timestamp_to_epoch_ms("2026-10-26 04:30:00.000")
        self.assertGreater(ms1, 1700000000000)

        # With T and Z
        ms2 = parse_timestamp_to_epoch_ms("2026-10-26T04:30:00.000Z")
        self.assertEqual(ms1, ms2)

        # Empty or invalid
        self.assertEqual(parse_timestamp_to_epoch_ms(""), 0.0)
        self.assertEqual(parse_timestamp_to_epoch_ms("invalid"), 0.0)

    def test_build_trace_trees_hierarchy(self):
        rows = [
            {
                "trace_id": "trace-1",
                "span_id": "root-1",
                "parent_span_id": "",
                "service_name": "frontend",
                "operation_name": "GET /home",
                "timestamp": "2026-10-26 00:00:00.000",
                "duration_ms": 100.0,
                "http_status": "200",
                "status_code": "OK",
                "kind": "SERVER",
            },
            {
                "trace_id": "trace-1",
                "span_id": "child-1",
                "parent_span_id": "root-1",
                "service_name": "backend",
                "operation_name": "QueryDB",
                "timestamp": "2026-10-26 00:00:00.020",
                "duration_ms": 40.0,
                "http_status": "200",
                "status_code": "OK",
                "kind": "CLIENT",
            },
        ]
        trees = build_trace_trees(rows)
        self.assertIn("trace-1", trees)
        t = trees["trace-1"]
        self.assertEqual(t["root_service"], "frontend")
        self.assertEqual(t["root_op"], "GET /home")
        self.assertEqual(t["spans_count"], 2)
        self.assertEqual(t["status"], "OK")

        # Check ordered spans and depth
        ordered = t["ordered_spans"]
        self.assertEqual(len(ordered), 2)
        self.assertEqual(ordered[0].span_id, "root-1")
        self.assertEqual(ordered[0].depth, 0)
        self.assertEqual(ordered[1].span_id, "child-1")
        self.assertEqual(ordered[1].depth, 1)
        self.assertIn("└── ", ordered[1].tree_prefix)
        self.assertEqual(ordered[1].offset_ms, 20.0)

    def test_build_trace_trees_error_status(self):
        rows = [
            {
                "trace_id": "trace-err",
                "span_id": "root-err",
                "parent_span_id": "",
                "service_name": "svc",
                "operation_name": "op",
                "timestamp": "2026-10-26 00:00:00.000",
                "duration_ms": 50.0,
                "http_status": "500",
                "status_code": "ERROR",
                "kind": "SERVER",
            }
        ]
        trees = build_trace_trees(rows)
        self.assertEqual(trees["trace-err"]["status"], "FAULT")
        self.assertEqual(trees["trace-err"]["error_count"], 1)

    def test_inject_report_end_to_end(self):
        # Create a sample CSV
        csv_file = tempfile.NamedTemporaryFile(delete=False, suffix=".csv", mode="w", newline="", encoding="utf-8")
        csv_path = csv_file.name
        csv_file.write(
            "timestamp,trace_id,span_id,parent_span_id,service_name,operation_name,duration_ms,duration_nano,http_status,http_method,http_url,status_code,status_message,kind\n"
            "2026-10-26 01:00:00.000,t-test-1,s-root,,frontend,GET /home,150.0,150000000,200,GET,/home,OK,,SERVER\n"
            "2026-10-26 01:00:00.020,t-test-1,s-child,s-root,db-svc,SELECT items,80.0,80000000,,,db:items,OK,,CLIENT\n"
            "2026-10-26 01:00:01.000,t-test-2,s-err,,auth-svc,POST /login,45.0,45000000,401,POST,/login,ERROR,Unauthorized,SERVER\n"
        )
        csv_file.close()

        out_excel = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        out_excel_path = out_excel.name
        out_excel.close()

        try:
            inject_report(
                csv_path=csv_path,
                template_path=self.template_path,
                output_path=out_excel_path,
                start_time="2026-10-26 01:00:00 UTC",
                end_time="2026-10-26 01:05:00 UTC",
                filter_criteria="All Transactions",
            )

            # Load and verify generated Excel
            wb = openpyxl.load_workbook(out_excel_path)
            self.assertIn("Summary", wb.sheetnames)
            self.assertIn("Trace_Waterfall", wb.sheetnames)
            self.assertIn("Raw_Data", wb.sheetnames)
            self.assertIn("Trace_2", wb.sheetnames)

            # Summary sheet validations
            ws_sum = wb["Summary"]
            self.assertEqual(ws_sum["B8"].value, 2)  # Total Traces
            self.assertEqual(ws_sum["E8"].value, 3)  # Total Spans
            self.assertEqual(ws_sum["H8"].value, 1)  # OK Traces
            self.assertEqual(ws_sum["J8"].value, 1)  # Error Traces

            # Waterfall sheet validations
            ws_wat = wb["Trace_Waterfall"]
            self.assertIsNotNone(ws_wat["B6"].value)
            self.assertGreaterEqual(len(ws_wat._charts), 1)

            # Raw_Data sheet validations
            ws_raw = wb["Raw_Data"]
            self.assertEqual(ws_raw.cell(row=5, column=2).value, "t-test-1")
            self.assertEqual(ws_raw.cell(row=6, column=2).value, "t-test-1")
            self.assertEqual(ws_raw.cell(row=7, column=2).value, "t-test-2")

        finally:
            if os.path.exists(csv_path):
                os.remove(csv_path)
            if os.path.exists(out_excel_path):
                os.remove(out_excel_path)

    def test_inject_report_empty_data(self):
        # Empty CSV
        csv_file = tempfile.NamedTemporaryFile(delete=False, suffix=".csv", mode="w", newline="", encoding="utf-8")
        csv_path = csv_file.name
        csv_file.write(
            "timestamp,trace_id,span_id,parent_span_id,service_name,operation_name,duration_ms,duration_nano,http_status,http_method,http_url,status_code,status_message,kind\n"
        )
        csv_file.close()

        out_excel = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        out_excel_path = out_excel.name
        out_excel.close()

        try:
            inject_report(
                csv_path=csv_path,
                template_path=self.template_path,
                output_path=out_excel_path,
            )
            wb = openpyxl.load_workbook(out_excel_path)
            ws_sum = wb["Summary"]
            self.assertEqual(ws_sum["B8"].value, 0)
            self.assertEqual(ws_sum["E8"].value, 0)
        finally:
            if os.path.exists(csv_path):
                os.remove(csv_path)
            if os.path.exists(out_excel_path):
                os.remove(out_excel_path)


if __name__ == "__main__":
    unittest.main()
