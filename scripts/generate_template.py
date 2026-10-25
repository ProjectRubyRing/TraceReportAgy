#!/usr/bin/env python3
"""
generate_template.py
Generate a professional Excel template (trace_report_template.xlsx)
modeled after the AWS CloudWatch / X-Ray trace details console.
Designed for RHEL 9.8 / Python 3.9+ with openpyxl.
"""

import sys
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Colors (Hex)
NAVY_HEADER = "1B2A4A"       # AWS Navy
WHITE = "FFFFFF"
GRAY_TEXT = "5A6A85"
CARD_BG = "F8FAFC"
CARD_BORDER = "E2E8F0"
BORDER_LIGHT = "E5E7EB"
BORDER_DARK = "CBD5E1"
SUCCESS_BG = "D1E7DD"
SUCCESS_FG = "0F5132"
WARNING_BG = "FFF3CD"
WARNING_FG = "664D03"
DANGER_BG = "F8D7DA"
DANGER_FG = "842029"
INFO_BG = "CFE2FF"
INFO_FG = "084298"
TIMELINE_BG = "F1F5F9"
ROW_ALT = "FBFDFF"

FONT_FAMILY = "Segoe UI"


def create_thin_border(color=BORDER_LIGHT):
    side = Side(border_style="thin", color=color)
    return Border(left=side, right=side, top=side, bottom=side)


def create_template(output_path: str = "templates/trace_report_template.xlsx"):
    wb = openpyxl.Workbook()

    # ----------------------------------------------------
    # Sheet 1: Summary (Trace Execution Index & Overview)
    # ----------------------------------------------------
    ws_sum = wb.active
    ws_sum.title = "Summary"
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner (Row 2 to 3, Cols B to M)
    ws_sum.merge_cells("B2:M3")
    title_cell = ws_sum["B2"]
    title_cell.value = "AWS CloudWatch Transaction Search - Trace Report"
    title_cell.font = Font(name=FONT_FAMILY, size=16, bold=True, color=WHITE)
    title_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Subtitle / Metadata (Row 4, 5)
    ws_sum["B4"] = "Log Group:"
    ws_sum["B4"].font = Font(name=FONT_FAMILY, size=10, bold=True, color=GRAY_TEXT)
    ws_sum["C4"] = "aws/spans"
    ws_sum["C4"].font = Font(name=FONT_FAMILY, size=10, bold=False)

    ws_sum["F4"] = "Query Window (UTC):"
    ws_sum["F4"].font = Font(name=FONT_FAMILY, size=10, bold=True, color=GRAY_TEXT)
    ws_sum["G4"] = "{{QUERY_START}} ~ {{QUERY_END}}"
    ws_sum["G4"].font = Font(name=FONT_FAMILY, size=10, bold=False)

    ws_sum["B5"] = "Generated At:"
    ws_sum["B5"].font = Font(name=FONT_FAMILY, size=10, bold=True, color=GRAY_TEXT)
    ws_sum["C5"] = "{{GENERATED_AT}}"
    ws_sum["C5"].font = Font(name=FONT_FAMILY, size=10, bold=False)

    ws_sum["F5"] = "Status Filter:"
    ws_sum["F5"].font = Font(name=FONT_FAMILY, size=10, bold=True, color=GRAY_TEXT)
    ws_sum["G5"] = "{{FILTER_CRITERIA}}"
    ws_sum["G5"].font = Font(name=FONT_FAMILY, size=10, bold=False)

    # KPI Summary Cards (Row 7 to 9)
    # Card 1: Total Traces
    # Card 2: Total Spans
    # Card 3: Success Rate / Count
    # Card 4: Error / Fault Count
    # Card 5: Max Latency
    kpi_configs = [
        ("B", "D", "TOTAL TRACES", "{{TOTAL_TRACES}}", CARD_BG, GRAY_TEXT),
        ("E", "G", "TOTAL SPANS", "{{TOTAL_SPANS}}", CARD_BG, GRAY_TEXT),
        ("H", "I", "OK (2xx)", "{{OK_TRACES}}", SUCCESS_BG, SUCCESS_FG),
        ("J", "K", "ERROR / FAULT", "{{ERROR_TRACES}}", DANGER_BG, DANGER_FG),
        ("L", "M", "MAX DURATION", "{{MAX_DURATION}}", CARD_BG, GRAY_TEXT),
    ]

    for start_col, end_col, label, val_placeholder, bg_color, text_color in kpi_configs:
        cell_lbl = f"{start_col}7"
        cell_val = f"{start_col}8"
        ws_sum.merge_cells(f"{start_col}7:{end_col}7")
        ws_sum.merge_cells(f"{start_col}8:{end_col}9")

        c_lbl = ws_sum[cell_lbl]
        c_lbl.value = label
        c_lbl.font = Font(name=FONT_FAMILY, size=9, bold=True, color=text_color)
        c_lbl.alignment = Alignment(horizontal="center", vertical="center")
        c_lbl.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        c_val = ws_sum[cell_val]
        c_val.value = val_placeholder
        c_val.font = Font(name=FONT_FAMILY, size=16, bold=True, color=text_color)
        c_val.alignment = Alignment(horizontal="center", vertical="center")
        c_val.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        # Set borders around card
        for row in range(7, 10):
            for col_letter in [chr(c) for c in range(ord(start_col), ord(end_col) + 1)]:
                ws_sum[f"{col_letter}{row}"].border = create_thin_border(CARD_BORDER)

    # Section Title for Table
    ws_sum["B11"] = "Trace Summary Table"
    ws_sum["B11"].font = Font(name=FONT_FAMILY, size=12, bold=True, color=NAVY_HEADER)

    # Trace Table Headers (Row 12)
    sum_headers = [
        ("Trace ID", 34),
        ("Start Time", 22),
        ("Root Service", 20),
        ("Root Operation / URL", 30),
        ("Duration (ms)", 16),
        ("Spans", 10),
        ("Services", 10),
        ("Status", 14),
        ("Errors", 10),
        ("Details Link", 16),
    ]

    for col_idx, (header_name, col_width) in enumerate(sum_headers, start=2):
        col_letter = get_column_letter(col_idx)
        cell = ws_sum.cell(row=12, column=col_idx, value=header_name)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = create_thin_border(BORDER_DARK)
        ws_sum.column_dimensions[col_letter].width = col_width

    ws_sum.row_dimensions[2].height = 24
    ws_sum.row_dimensions[3].height = 20
    ws_sum.row_dimensions[7].height = 18
    ws_sum.row_dimensions[8].height = 22
    ws_sum.row_dimensions[9].height = 18
    ws_sum.row_dimensions[12].height = 24
    ws_sum.column_dimensions["A"].width = 3

    # ----------------------------------------------------
    # Sheet 2: Trace_Waterfall (X-Ray Timeline View)
    # ----------------------------------------------------
    ws_wat = wb.create_sheet(title="Trace_Waterfall")
    ws_wat.views.sheetView[0].showGridLines = True

    # Title Banner (Row 2-3)
    ws_wat.merge_cells("B2:U3")
    t_cell = ws_wat["B2"]
    t_cell.value = "AWS X-Ray Trace Timeline & Waterfall View"
    t_cell.font = Font(name=FONT_FAMILY, size=16, bold=True, color=WHITE)
    t_cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    t_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    # Trace KPI Overview Bar (Row 5 to 7)
    # Card A: Trace ID
    # Card B: Status
    # Card C: Root Service & Operation
    # Card D: Total Duration
    # Card E: Total Spans
    # Card F: Start Time
    trace_kpis = [
        ("B", "E", "TRACE ID", "{{CURRENT_TRACE_ID}}", CARD_BG, GRAY_TEXT, 11),
        ("F", "G", "STATUS", "{{CURRENT_STATUS}}", SUCCESS_BG, SUCCESS_FG, 14),
        ("H", "L", "ROOT SERVICE / OPERATION", "{{CURRENT_ROOT}}", CARD_BG, GRAY_TEXT, 11),
        ("M", "N", "TOTAL DURATION", "{{CURRENT_DURATION}}", CARD_BG, GRAY_TEXT, 14),
        ("O", "P", "TOTAL SPANS", "{{CURRENT_SPANS_COUNT}}", CARD_BG, GRAY_TEXT, 14),
        ("Q", "U", "START TIMESTAMP (UTC)", "{{CURRENT_START_TIME}}", CARD_BG, GRAY_TEXT, 11),
    ]

    for start_col, end_col, label, val_placeholder, bg_color, text_color, val_size in trace_kpis:
        ws_wat.merge_cells(f"{start_col}5:{end_col}5")
        ws_wat.merge_cells(f"{start_col}6:{end_col}7")

        lbl_c = ws_wat[f"{start_col}5"]
        lbl_c.value = label
        lbl_c.font = Font(name=FONT_FAMILY, size=8, bold=True, color=text_color)
        lbl_c.alignment = Alignment(horizontal="center", vertical="center")
        lbl_c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        val_c = ws_wat[f"{start_col}6"]
        val_c.value = val_placeholder
        val_c.font = Font(name=FONT_FAMILY, size=val_size, bold=True, color=text_color)
        val_c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        val_c.fill = PatternFill(start_color=bg_color, end_color=bg_color, fill_type="solid")

        for row in range(5, 8):
            for col_letter in [chr(c) for c in range(ord(start_col[0]), ord(end_col[0]) + 1)]:
                ws_wat[f"{col_letter}{row}"].border = create_thin_border(CARD_BORDER)

    # Section Title
    ws_wat["B9"] = "Timeline Segments & Waterfall"
    ws_wat["B9"].font = Font(name=FONT_FAMILY, size=12, bold=True, color=NAVY_HEADER)

    # Timeline Table Headers (Row 10)
    # Left columns: Span attributes (B to I)
    # Right columns: Timeline bar visualization (J to AC: 20 buckets)
    span_headers = [
        ("Status", 10),             # B
        ("Hierarchy & Operation", 40), # C
        ("Service", 18),            # D
        ("Kind", 12),               # E
        ("HTTP", 8),                # F
        ("Start Offset", 14),       # G
        ("Duration (ms)", 14),      # H
        ("% of Trace", 12),         # I
    ]

    for col_idx, (header_name, col_width) in enumerate(span_headers, start=2):
        col_letter = get_column_letter(col_idx)
        cell = ws_wat.cell(row=10, column=col_idx, value=header_name)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = create_thin_border(BORDER_DARK)
        ws_wat.column_dimensions[col_letter].width = col_width

    # Timeline Bar Buckets (Cols J to AC, 20 intervals: 0% to 100%)
    timeline_cols_count = 20
    timeline_start_col = 10  # Col J

    ws_wat.merge_cells(start_row=9, start_column=timeline_start_col, end_row=9, end_column=timeline_start_col + timeline_cols_count - 1)
    timeline_label = ws_wat.cell(row=9, column=timeline_start_col, value="Waterfall Timeline (0% ------------------------> 100%)")
    timeline_label.font = Font(name=FONT_FAMILY, size=9, bold=True, color=GRAY_TEXT)
    timeline_label.alignment = Alignment(horizontal="center", vertical="center")

    for i in range(timeline_cols_count):
        col_idx = timeline_start_col + i
        col_letter = get_column_letter(col_idx)
        pct_label = f"{int(i * 100 / timeline_cols_count)}%"
        cell = ws_wat.cell(row=10, column=col_idx, value=pct_label)
        cell.font = Font(name=FONT_FAMILY, size=8, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color="334E68", end_color="334E68", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = create_thin_border(BORDER_DARK)
        ws_wat.column_dimensions[col_letter].width = 4.0

    ws_wat.row_dimensions[2].height = 24
    ws_wat.row_dimensions[3].height = 20
    ws_wat.row_dimensions[5].height = 18
    ws_wat.row_dimensions[6].height = 22
    ws_wat.row_dimensions[7].height = 18
    ws_wat.row_dimensions[10].height = 24
    ws_wat.column_dimensions["A"].width = 3

    # ----------------------------------------------------
    # Sheet 3: Raw_Data (Spans Flat Data)
    # ----------------------------------------------------
    ws_raw = wb.create_sheet(title="Raw_Data")
    ws_raw.views.sheetView[0].showGridLines = True

    # Title Banner (Row 2)
    ws_raw.merge_cells("A2:N2")
    raw_title = ws_raw["A2"]
    raw_title.value = "CloudWatch Logs Insights - Exported Spans (aws/spans)"
    raw_title.font = Font(name=FONT_FAMILY, size=14, bold=True, color=WHITE)
    raw_title.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
    raw_title.alignment = Alignment(horizontal="left", vertical="center", indent=1)

    raw_headers = [
        ("timestamp", 22),
        ("trace_id", 34),
        ("span_id", 18),
        ("parent_span_id", 18),
        ("service_name", 20),
        ("operation_name", 30),
        ("duration_ms", 14),
        ("duration_nano", 16),
        ("http_status", 12),
        ("http_method", 12),
        ("http_url", 35),
        ("status_code", 14),
        ("status_message", 30),
        ("kind", 14),
    ]

    for col_idx, (h_name, h_width) in enumerate(raw_headers, start=1):
        col_letter = get_column_letter(col_idx)
        cell = ws_raw.cell(row=4, column=col_idx, value=h_name)
        cell.font = Font(name=FONT_FAMILY, size=10, bold=True, color=WHITE)
        cell.fill = PatternFill(start_color=NAVY_HEADER, end_color=NAVY_HEADER, fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = create_thin_border(BORDER_DARK)
        ws_raw.column_dimensions[col_letter].width = h_width

    ws_raw.row_dimensions[2].height = 28
    ws_raw.row_dimensions[4].height = 22

    # Save template
    wb.save(output_path)
    print(f"Created Excel template successfully: {output_path}")


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "templates/trace_report_template.xlsx"
    create_template(out)
